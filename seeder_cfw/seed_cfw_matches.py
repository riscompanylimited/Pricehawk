#!/usr/bin/env python3
"""
Seed CFW<->Makro match pairs into `product_matches` from the survey
(Example_Survey_(FB1FB2FB3FE)-2.xlsx).

Each survey data row carries a CFW SKU (col C 'SKU'/'Item') and its matched
Makro SKU (col B 'SKU Makro'). This maps both to products.id and upserts a row
into product_matches.

Since the survey is a curated official match list, matches are written as
CONFIRMED: is_verified=TRUE, is_same=TRUE, verified_at=NOW() — which is exactly
what the comparison UI and update_makro_location_prices.py require. Pairs whose
Makro SKU was never seeded (the 12 delisted ones) are skipped.

Usage:
    python seed_cfw_matches.py            # write
    python seed_cfw_matches.py --dry-run  # summarise, no write
"""
import argparse
import sys
from pathlib import Path
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
REPO = HERE.parent
SURVEY_XLSX = HERE / "Example_Survey_(FB1FB2FB3FE)-2.xlsx"
ENV_PATH = REPO / "backend" / ".env"


def load_env():
    env = {}
    if ENV_PATH.exists():
        for line in ENV_PATH.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if line and not line.startswith("#") and "=" in line:
                k, _, v = line.partition("=")
                env[k.strip()] = v.strip()
    return env


def get_conn(env):
    import psycopg2
    from psycopg2.extras import RealDictCursor
    return psycopg2.connect(
        host=env.get("DB_HOST", "localhost"), port=int(env.get("DB_PORT", 5432)),
        dbname=env.get("DB_NAME", "pricehawk"), user=env.get("DB_USER", "postgres"),
        password=env.get("DB_PASSWORD", ""), sslmode=env.get("DB_SSLMODE", "prefer"),
        cursor_factory=RealDictCursor,
    )


def _sku(v):
    if v is None:
        return None
    s = str(v).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return s or None


def read_pairs():
    """(cfw_sku, makro_sku) from every survey sheet; deduped, order-preserving."""
    wb = load_workbook(SURVEY_XLSX, read_only=True, data_only=True)
    seen, pairs = set(), []
    for ws in wb.worksheets:
        for r in ws.iter_rows(values_only=True):
            a = r[0] if len(r) > 0 else None     # 'No.' -> int on data rows
            mk = _sku(r[1]) if len(r) > 1 else None   # SKU Makro
            cf = _sku(r[2]) if len(r) > 2 else None   # CFW SKU
            if isinstance(a, (int, float)) and mk and cf and mk.isdigit() and cf.isdigit():
                key = (cf, mk)
                if key not in seen:
                    seen.add(key)
                    pairs.append(key)
    wb.close()
    return pairs


def main():
    ap = argparse.ArgumentParser(description="Seed CFW<->Makro matches from the survey.")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    if not SURVEY_XLSX.exists():
        print(f"[ERROR] missing {SURVEY_XLSX}"); return 1

    pairs = read_pairs()
    print(f"[INFO] survey pairs (unique cfw+makro): {len(pairs)}")

    env = load_env()
    conn = get_conn(env)
    cur = conn.cursor()
    cur.execute("SELECT sku, id FROM products WHERE retailer_id='cfw'")
    cfw_id = {r["sku"]: r["id"] for r in cur.fetchall()}
    cur.execute("SELECT sku, id FROM products WHERE retailer_id='makro'")
    makro_id = {r["sku"]: r["id"] for r in cur.fetchall()}

    resolved, miss_cfw, miss_makro = [], [], []
    for cf, mk in pairs:
        c, m = cfw_id.get(cf), makro_id.get(mk)
        if c is None:
            miss_cfw.append(cf)
        elif m is None:
            miss_makro.append(mk)
        else:
            resolved.append((c, m, cf, mk))

    print(f"[INFO] resolvable pairs: {len(resolved)}")
    if miss_cfw:
        print(f"[WARN] {len(miss_cfw)} pairs skipped — CFW SKU not in products: {sorted(set(miss_cfw))[:8]}")
    if miss_makro:
        print(f"[WARN] {len(miss_makro)} pairs skipped — Makro SKU not in products (delisted): {sorted(set(miss_makro))}")

    if args.dry_run:
        print("[DRY RUN] no write.")
        for c, m, cf, mk in resolved[:5]:
            print(f"  cfw {cf}(#{c}) <-> makro {mk}(#{m})")
        conn.close(); return 0

    ins = 0
    for c, m, cf, mk in resolved:
        cur.execute("""
            INSERT INTO product_matches
                (cfw_product_id, makro_product_id, match_score, is_verified, is_same, verified_at, created_at, updated_at)
            VALUES (%s, %s, NULL, TRUE, TRUE, NOW(), NOW(), NOW())
            ON CONFLICT (cfw_product_id, makro_product_id) DO UPDATE SET
                is_verified=TRUE, is_same=TRUE, verified_at=NOW(), updated_at=NOW()
        """, (c, m))
        ins += 1
    conn.commit()
    cur.execute("SELECT count(*) AS n FROM product_matches WHERE is_verified AND is_same")
    total = cur.fetchone()["n"]
    print(f"[OK] upserted {ins} matches; product_matches now has {total} verified rows")
    cur.close(); conn.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
