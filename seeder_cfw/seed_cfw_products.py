#!/usr/bin/env python3
"""
Seed the 159 matched CFW products into `products`.

Scope (grilled 2026-07-12): insert ONLY the CFW SKUs that appear in the
CFW<->Makro match survey (Example_Survey_(FB1FB2FB3FE)-2.xlsx), pulling their
full attributes from the CFW master template (Product_template_CFW.xlsx).
Does NOT touch Makro rows, product_matches, or watchlists — those are later steps.

Decisions:
- SKU set        : the survey's CFW SKUs (col C 'SKU'/'Item'), all present in master.
- Hierarchy      : seed the distinct categories/departments/classes for those SKUs
                   first (products has FKs on category/dept/class), then products.
- step_prices    : [] for all (only 3 SKUs have promo data; deferred).
- current_price  : POS PRICE from the master 'Price' sheet (STCODE 1).
- is_active      : SKU Status 'A-Active' -> TRUE, else FALSE.
- On conflict    : full overwrite from Excel (Excel = source of truth).

Usage:
    python seed_cfw_products.py            # write to DB
    python seed_cfw_products.py --dry-run  # parse + summarise, no DB write

Reads DB creds from backend/.env (DB_HOST/DB_PORT/DB_NAME/DB_USER/DB_PASSWORD/DB_SSLMODE).
"""
import argparse
import json
import os
import sys
from pathlib import Path
from collections import Counter
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
REPO = HERE.parent
MASTER_XLSX = HERE / "Product_template_CFW.xlsx"
SURVEY_XLSX = HERE / "Example_Survey_(FB1FB2FB3FE)-2.xlsx"
ENV_PATH = REPO / "backend" / ".env"

NA = "N/A"  # placeholder for NULL sub_dept/sub_class (PK can't be NULL)


def load_env():
    env = {}
    if ENV_PATH.exists():
        for line in ENV_PATH.read_text(encoding="utf-8").splitlines():
            line = line.strip()
            if line and not line.startswith("#") and "=" in line:
                k, _, v = line.partition("=")
                env[k.strip()] = v.strip()
    return env


def get_conn(env):
    import psycopg2
    from psycopg2.extras import RealDictCursor
    return psycopg2.connect(
        host=env.get("DB_HOST", "localhost"),
        port=int(env.get("DB_PORT", 5432)),
        dbname=env.get("DB_NAME", "pricehawk"),
        user=env.get("DB_USER", "postgres"),
        password=env.get("DB_PASSWORD", ""),
        sslmode=env.get("DB_SSLMODE", "prefer"),
        cursor_factory=RealDictCursor,
    )


def _s(v):
    """Excel cell -> clean str or None."""
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _sku(v):
    """Normalise a SKU-ish cell to a plain digit string ('10005147.0' -> '10005147')."""
    if v is None:
        return None
    s = str(v).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return s or None


def read_survey_cfw_skus():
    """The CFW SKUs (col C) from every survey sheet — the matched set."""
    wb = load_workbook(SURVEY_XLSX, read_only=True, data_only=True)
    skus = set()
    for ws in wb.worksheets:
        for r in ws.iter_rows(values_only=True):
            a = r[0] if len(r) > 0 else None       # 'No.' -> int on data rows
            c = r[2] if len(r) > 2 else None        # CFW SKU
            if isinstance(a, (int, float)) and c is not None:
                s = _sku(c)
                if s and s.isdigit():
                    skus.add(s)
    wb.close()
    return skus


def read_price_map():
    """SKU -> POS PRICE from the master 'Price' sheet (STCODE 1)."""
    wb = load_workbook(MASTER_XLSX, read_only=True, data_only=True)
    pm = {}
    for r in wb["Price"].iter_rows(min_row=2, values_only=True):
        sku, price = _sku(r[1]) if len(r) > 1 else None, (r[2] if len(r) > 2 else None)
        if sku and price is not None:
            try:
                pm[sku] = float(price)
            except (TypeError, ValueError):
                pass
    wb.close()
    return pm


def read_master_rows(wanted):
    """Rows from the master Products sheet whose SKU is in `wanted`, keyed by header."""
    wb = load_workbook(MASTER_XLSX, read_only=True, data_only=True)
    ws = wb["Products"]
    rows_iter = ws.iter_rows(values_only=True)
    headers = [(_s(h) or f"col{i}") for i, h in enumerate(next(rows_iter))]
    out = {}
    for r in rows_iter:
        sku = _sku(r[0]) if r else None
        if sku and sku in wanted:
            out[sku] = {headers[i]: r[i] for i in range(min(len(headers), len(r)))}
    wb.close()
    return out


def main():
    ap = argparse.ArgumentParser(description="Seed the matched CFW products.")
    ap.add_argument("--dry-run", action="store_true", help="parse + summarise, no DB write")
    args = ap.parse_args()

    for f in (MASTER_XLSX, SURVEY_XLSX):
        if not f.exists():
            print(f"[ERROR] missing {f}"); return 1

    wanted = read_survey_cfw_skus()
    print(f"[INFO] survey CFW SKUs: {len(wanted)}")
    price_map = read_price_map()
    master = read_master_rows(wanted)
    print(f"[INFO] matched in master: {len(master)}")

    missing = wanted - set(master)
    if missing:
        print(f"[WARN] {len(missing)} survey SKUs not in master: {sorted(missing)[:10]}")

    # Build product rows + distinct parents
    categories = {}                 # (cfw, cat_id) -> cat_name
    departments = {}                # (dept_id, sub_dept_id) -> (dept_name, sub_dept_name)
    classes = {}                    # (class_id, sub_class_id) -> (class_name, sub_class_name)
    products = []
    status_counter = Counter()

    for sku, row in master.items():
        cat_id = _s(row.get("Category ID"))
        cat_name = _s(row.get("Category Name"))
        dept_id = _s(row.get("Dept"))
        dept_name = _s(row.get("Dept Name"))
        sub_dept_id = _s(row.get("Sub-Dept")) or NA
        sub_dept_name = _s(row.get("Sub-Dept Name")) or NA
        class_id = _s(row.get("Class"))
        class_name = _s(row.get("Class Name"))
        sub_class_id = _s(row.get("Sub-Class")) or NA
        sub_class_name = _s(row.get("Sub-Class Name")) or NA
        status = _s(row.get("SKU Status")) or ""
        is_active = status.startswith("A")
        status_counter[status] += 1

        if cat_id and cat_name:
            categories[("cfw", cat_id)] = cat_name
        if dept_id and dept_name:
            departments[(dept_id, sub_dept_id)] = (dept_name, sub_dept_name)
        if class_id and class_name:
            classes[(class_id, sub_class_id)] = (class_name, sub_class_name)

        products.append({
            "sku": sku,
            "barcode": _s(row.get("Barcode")),
            "name": _s(row.get("Product Name (TH)")) or sku,
            "name_en": _s(row.get("Product Name (EN)")),
            "brand": _s(row.get("Brand")),
            "category_id": cat_id,
            "dept_id": dept_id,
            "sub_dept_id": sub_dept_id if dept_id else None,
            "class_id": class_id,
            "sub_class_id": sub_class_id if class_id else None,
            "current_price": price_map.get(sku),
            "url": _s(row.get("Product URL")),
            "image_url": _s(row.get("Image URL")),
            "is_active": is_active,
        })

    print(f"[INFO] parents: {len(categories)} categories, {len(departments)} dept-pairs, "
          f"{len(classes)} class-pairs")
    print(f"[INFO] status: {dict(status_counter)}  "
          f"(active={sum(1 for p in products if p['is_active'])})")
    print(f"[INFO] with price: {sum(1 for p in products if p['current_price'] is not None)}/{len(products)}")

    if args.dry_run:
        print("[DRY RUN] no DB write.")
        for p in products[:3]:
            print("  sample:", {k: p[k] for k in ("sku", "name", "current_price", "category_id", "is_active")})
        return 0

    env = load_env()
    conn = get_conn(env)
    print(f"[INFO] DB: {env.get('DB_NAME')} @ {env.get('DB_HOST')}")
    cur = conn.cursor()
    try:
        # 1. parents first (FK targets)
        for (rid, cid), name in categories.items():
            cur.execute("""INSERT INTO categories (retailer_id, category_id, category_name)
                           VALUES (%s,%s,%s)
                           ON CONFLICT (retailer_id, category_id)
                           DO UPDATE SET category_name = EXCLUDED.category_name""",
                        (rid, cid, name))
        for (did, sdid), (dn, sdn) in departments.items():
            cur.execute("""INSERT INTO departments (dept_id, dept_name, sub_dept_id, sub_dept_name)
                           VALUES (%s,%s,%s,%s)
                           ON CONFLICT (dept_id, sub_dept_id)
                           DO UPDATE SET dept_name=EXCLUDED.dept_name, sub_dept_name=EXCLUDED.sub_dept_name""",
                        (did, dn, sdid, sdn))
        for (cid, scid), (cn, scn) in classes.items():
            cur.execute("""INSERT INTO classes (class_id, class_name, sub_class_id, sub_class_name)
                           VALUES (%s,%s,%s,%s)
                           ON CONFLICT (class_id, sub_class_id)
                           DO UPDATE SET class_name=EXCLUDED.class_name, sub_class_name=EXCLUDED.sub_class_name""",
                        (cid, cn, scid, scn))
        conn.commit()
        print(f"[OK] seeded {len(categories)} categories, {len(departments)} departments, {len(classes)} classes")

        # 2. products (full overwrite on conflict)
        ins = 0
        for p in products:
            cur.execute("""
                INSERT INTO products (
                    retailer_id, sku, barcode, name, name_en, brand,
                    category_id, dept_id, sub_dept_id, class_id, sub_class_id,
                    current_price, step_prices, url, image_url, is_active, created_at, updated_at
                ) VALUES ('cfw', %s,%s,%s,%s,%s, %s,%s,%s,%s,%s, %s,%s,%s,%s,%s, NOW(), NOW())
                ON CONFLICT (retailer_id, sku) DO UPDATE SET
                    barcode=EXCLUDED.barcode, name=EXCLUDED.name, name_en=EXCLUDED.name_en,
                    brand=EXCLUDED.brand, category_id=EXCLUDED.category_id,
                    dept_id=EXCLUDED.dept_id, sub_dept_id=EXCLUDED.sub_dept_id,
                    class_id=EXCLUDED.class_id, sub_class_id=EXCLUDED.sub_class_id,
                    current_price=EXCLUDED.current_price, step_prices=EXCLUDED.step_prices,
                    url=EXCLUDED.url, image_url=EXCLUDED.image_url,
                    is_active=EXCLUDED.is_active, updated_at=NOW()
            """, (
                p["sku"], p["barcode"], p["name"], p["name_en"], p["brand"],
                p["category_id"], p["dept_id"], p["sub_dept_id"], p["class_id"], p["sub_class_id"],
                p["current_price"], json.dumps([]), p["url"], p["image_url"], p["is_active"],
            ))
            ins += 1
        conn.commit()
        print(f"[OK] upserted {ins} CFW products")
    except Exception as e:
        conn.rollback()
        print(f"[ERROR] rolled back: {e}")
        raise
    finally:
        cur.close(); conn.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
